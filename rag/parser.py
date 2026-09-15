"""
parser.py — Extract plain text from uploaded NSE filing files.

Supported formats:
    .pdf   → pdfplumber
    .docx  → python-docx
    .txt   → built-in
    .xlsx  → openpyxl
    .csv   → built-in csv reader

Each parser returns a plain-text string.  Empty or unreadable files
return an empty string so the caller can skip them gracefully.
"""

import csv
#import io
#import os
from typing import Optional, Any


# ──────────────────────────────────────────────────────────────────────────────
# Public API
# ──────────────────────────────────────────────────────────────────────────────

def extract_text(file_path: str, filename: Optional[str] = None) -> str:
    """
    Extract text content from a file at *file_path*.

    Args:
        file_path (str): Absolute path to the file on disk.
        filename  (str): Original filename (used to determine extension when
                         *file_path* has a temp name). Defaults to *file_path*.

    Returns:
        str: Extracted plain text (may be empty on parse failure).
    """
    name = (filename or file_path).lower()

    if name.endswith(".pdf"):
        return _parse_pdf(file_path)
    elif name.endswith(".docx"):
        return _parse_docx(file_path)
    elif name.endswith(".txt"):
        return _parse_txt(file_path)
    elif name.endswith(".xlsx"):
        return _parse_xlsx(file_path)
    elif name.endswith(".csv"):
        return _parse_csv(file_path)
    else:
        return ""


# ──────────────────────────────────────────────────────────────────────────────
# Internal parsers
# ──────────────────────────────────────────────────────────────────────────────

def _parse_pdf(path: str) -> str:
    """Extract text from a PDF using pdfplumber."""
    try:
        import pdfplumber  # lazy import so missing dep gives a clear error
        parts:list[Any] = []
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                text = page.extract_text()
                if text:
                    parts.append(text.strip())
        return "\n\n".join(parts)
    except ImportError:
        raise RuntimeError(
            "pdfplumber is not installed. Run: pip install pdfplumber"
        )
    except Exception as e:
        print(f"[parser] PDF parse error for {path}: {e}")
        return ""


def _parse_docx(path: str) -> str:
    """Extract text from a .docx file using python-docx."""
    try:
        from docx import Document  # python-docx exposes as 'docx'
        doc = Document(path)
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(
                    cell.text.strip() for cell in row.cells if cell.text.strip()
                )
                if row_text:
                    paragraphs.append(row_text)
        return "\n\n".join(paragraphs)
    except ImportError:
        raise RuntimeError(
            "python-docx is not installed. Run: pip install python-docx"
        )
    except Exception as e:
        print(f"[parser] DOCX parse error for {path}: {e}")
        return ""


def _parse_txt(path: str) -> str:
    """Read a plain-text file, trying UTF-8 then latin-1 as fallback."""
    for encoding in ("utf-8", "latin-1", "cp1252"):
        try:
            with open(path, "r", encoding=encoding) as fh:
                return fh.read()
        except UnicodeDecodeError:
            continue
        except Exception as e:
            print(f"[parser] TXT parse error for {path}: {e}")
            return ""
    return ""


def _parse_xlsx(path: str) -> str:
    """Extract text from an Excel file using openpyxl (read-only mode)."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        parts:list[Any] = []
        for sheet in wb.worksheets:
            sheet_rows:list[Any] = []
            for row in sheet.iter_rows(values_only=True):
                row_text = " | ".join(
                    str(cell) for cell in row if cell is not None and str(cell).strip()
                )
                if row_text:
                    sheet_rows.append(row_text)
            if sheet_rows:
                parts.append(f"[Sheet: {sheet.title}]\n" + "\n".join(sheet_rows))
        wb.close()
        return "\n\n".join(parts)
    except ImportError:
        raise RuntimeError(
            "openpyxl is not installed. Run: pip install openpyxl"
        )
    except Exception as e:
        print(f"[parser] XLSX parse error for {path}: {e}")
        return ""


def _parse_csv(path: str) -> str:
    """Read a CSV file and convert to a pipe-delimited text block."""
    for encoding in ("utf-8", "latin-1", "cp1252"):
        try:
            with open(path, "r", encoding=encoding, newline="") as fh:
                reader = csv.reader(fh)
                rows = [" | ".join(cell.strip() for cell in row) for row in reader if any(row)]
            return "\n".join(rows)
        except UnicodeDecodeError:
            continue
        except Exception as e:
            print(f"[parser] CSV parse error for {path}: {e}")
            return ""
    return ""
